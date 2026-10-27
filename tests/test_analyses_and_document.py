import unittest
from pathlib import Path

from docx import Document

import tests  # noqa: F401
import src.JBGtranscriber as J
from tests.support import FakeOpenAI, completion, in_temp_dir, make_transcriber

DIALOGUE = ("Ordförande: Välkomna till mötet.\nDeltagare 1: Tack.\n\n### TALARREGISTER\n"
            "Ordförande: leder mötet\nDeltagare 1: deltar")
ORIGINAL = "Och då fick vi lufflarlagen att ta hänsyn till."
MARKED = "Och då fick vi [FEL?]lufflarlagen (Kanske: LAS-lagen)[/FEL?] att ta hänsyn till."


def answer(system, user):
    return DIALOGUE if "SVARETS FORM" in system else "Resultat."


def run(fake, transcription_format="raw", marked=None, **steps):
    t = make_transcriber(ORIGINAL)
    if marked is not None:
        def fake_marking():
            t.current_task = "Markering av misstänkta fel"
            t.marked_text = marked
        t.find_suspicious_phrases = fake_marking
    t.transcribe = lambda: None
    t.export_path = Path("out.docx")
    with fake.installed():
        t.perform_transcription_steps(transcription_format=transcription_format, **steps)
    document = Document("out.docx")
    headings = [p.text for p in document.paragraphs if p.style.name.startswith("Heading")]
    return t, document, headings


class InstructionTest(unittest.TestCase):
    def test_custom_analysis_gets_fixed_opening(self):
        fake = FakeOpenAI(answer)
        t = make_transcriber()
        with fake.installed():
            t.run_custom_analysis("Lista alla beslut.")
        opening = J.JBGtranscriber.analysis_prompt_defaults()["custom_analysis_preamble"]
        self.assertEqual(fake.requests[0]["system"], opening + "\nLista alla beslut.")
        self.assertEqual(t.model_calls[0]["task"], "Anpassad analys")

    def test_follow_up_standard_and_edited(self):
        fake = FakeOpenAI(answer)
        with fake.installed():
            make_transcriber().suggest_follow_up_questions()
            make_transcriber().suggest_follow_up_questions(custom_prompt="Tre frågor.")
        standard = J.JBGtranscriber.analysis_prompt_defaults()["follow_up_questions"]
        self.assertEqual([r["system"] for r in fake.requests], [standard, "Tre frågor."])
        self.assertTrue(standard.endswith("."))

    def test_speaker_format_is_added_by_the_code(self):
        fake = FakeOpenAI(answer)
        t = make_transcriber()
        with fake.installed():
            t.do_analyze_speakers_splitted(custom_prompt="Ett möte.")
        self.assertTrue(fake.requests[0]["system"].startswith("Ett möte."))
        self.assertIn("SVARETS FORM", fake.requests[0]["system"])
        self.assertIn("Ordförande: Välkomna", t.analyze_speakers)


class PresentationTest(unittest.TestCase):
    def test_speaker_presentation_with_bold_labels(self):
        with in_temp_dir():
            t, document, headings = run(FakeOpenAI(answer), "speakers")
            self.assertEqual(headings[0], "Transkribering med talarangivelser")
            start = [p.text for p in document.paragraphs].index(headings[0]) + 1
            bold = ["".join(r.text for r in p.runs if r.bold) for p in document.paragraphs[start:start + 2]]
            self.assertEqual(bold, ["Ordförande: ", "Deltagare 1: "])

    def test_failed_speaker_identification_falls_back(self):
        def failing(system, user):
            if "SVARETS FORM" in system:
                raise RuntimeError("API down")
            return "Resultat."
        with in_temp_dir():
            t, document, headings = run(FakeOpenAI(failing), "speakers")
            texts = [p.text for p in document.paragraphs]
            self.assertEqual(headings[0], "Transkribering")
            self.assertIn("[Talaridentifieringen kunde inte genomföras. Rå text visas i stället.]", texts)

    def test_analyses_use_marked_text_only_when_marked(self):
        with in_temp_dir():
            fake = FakeOpenAI(answer)
            t, document, _ = run(fake, "marked", marked=MARKED, generate_summary=True,
                                 suggest_follow_up_questions=True, custom_analysis_prompt="Lista beslut.")
            for request in fake.requests:
                self.assertIn(MARKED, request["user"])
                self.assertTrue(request["system"].endswith(J.MARKED_TEXT_NOTE))
            note = [p.text for p in document.paragraphs if "bygger på texten med markerade" in p.text]
            self.assertTrue(note[0].startswith("Sammanfattningen, följdfrågorna och den anpassade analysen"))

            fake = FakeOpenAI(answer)
            run(fake, "raw", generate_summary=True)
            self.assertNotIn("OM TEXTENS MARKERINGAR", fake.requests[0]["system"])

    def test_sections_order_statistics_and_cost(self):
        with in_temp_dir():
            fake = FakeOpenAI(lambda s, u: completion(answer(s, u), prompt_tokens=2281, completion_tokens=1653))
            t, document, headings = run(fake, "raw", generate_summary=True,
                                        suggest_follow_up_questions=True, custom_analysis_prompt="Lista beslut.")
            self.assertEqual(headings, ["Transkribering", "Sammanfattning", "Uppföljningsfrågor",
                                        "Anpassad analys", "Om transkriberingen", "Statistik"])
            table = document.tables[0]
            self.assertEqual(table.rows[0].cells[-1].text, "Kostnad (USD)")
            self.assertEqual([r.cells[0].text for r in table.rows][1:],
                             ["Sammanfattning", "Uppföljningsfrågor", "Anpassad analys", "Totalt"])
            self.assertEqual(table.rows[1].cells[-1].text, "0,03")
            texts = [p.text for p in document.paragraphs]
            self.assertTrue(any(x.startswith("Uppskattad kostnad för OpenAI-anropen: 0,08 USD") for x in texts))
            self.assertTrue(any(x.startswith("Programversion: ") for x in texts))


class SpeakerLayoutTest(unittest.TestCase):
    def render(self, content):
        document = Document()
        J.JBGtranscriber._add_plain_word_content(document, content, parse_inline=True, speaker_labels=True)
        return [("".join(r.text for r in p.runs if r.bold), p.text) for p in document.paragraphs]

    def test_labels_continuations_and_colons_in_replies(self):
        content = J.JBGtranscriber._append_speaker_register(
            "Intervjuare 1: Hur?\nIntervjuobjekt 1: Det var så: vi gick.\nFortsättning.\n**Intervjuare 2**: Tack.",
            {"Intervjuare 1": "frågar", "Intervjuobjekt 1": "svarar"})
        rows = self.render(content)
        self.assertEqual([b for b, _ in rows[:3]], ["Intervjuare 1: ", "Intervjuobjekt 1: ", "Intervjuare 2: "])
        self.assertEqual(rows[1][1], "Intervjuobjekt 1: Det var så: vi gick.\nFortsättning.")
